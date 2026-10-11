import io

import docx
import pytest

from citewise.chunking import chunk_sections, split_text
from citewise.parsing import DOCX, MARKDOWN, PDF, TEXT, UnsupportedFile, detect_media_type, parse
from tests.conftest import make_pdf


def test_pdf_sections_keep_real_page_numbers():
    data = make_pdf(["Alpha revenue grew.", "Beta margins shrank.", "Gamma outlook stable."])
    parsed = parse(data, PDF)
    assert parsed.page_count == 3
    assert [s.page_number for s in parsed.sections] == [1, 2, 3]
    assert "Beta" in parsed.sections[1].text


def test_chunks_never_cross_page_boundaries():
    long_page = " ".join(f"Sentence {i} about page one." for i in range(120))
    data = make_pdf([long_page, "Short second page with the key figure 42."])
    chunks = list(chunk_sections(parse(data, PDF).sections, 500, 50, 120))
    assert {c.page_number for c in chunks} == {1, 2}
    second = [c for c in chunks if c.page_number == 2]
    assert len(second) == 1 and "42" in second[0].text
    assert all("page one" not in c.text for c in second)
    assert [c.index for c in chunks] == list(range(len(chunks)))
    assert all(len(c.text) <= 500 for c in chunks)
    assert all(c.children for c in chunks)


def test_split_text_respects_size_and_overlap():
    text = "\n\n".join(f"Paragraph {i} " + "word " * 40 for i in range(10))
    pieces = split_text(text, 300, 40)
    assert len(pieces) > 1
    assert all(len(p) <= 300 + 40 for p in pieces)
    joined = " ".join(pieces)
    assert all(f"Paragraph {i}" in joined for i in range(10))


def test_markdown_sections_by_heading_ignoring_code_fences():
    md = "# Intro\nHello\n```\n# not a heading\n```\n## Pricing\nCosts $10.\n"
    parsed = parse(md.encode(), MARKDOWN)
    headings = [s.heading for s in parsed.sections]
    assert headings == ["Intro", "Pricing"]
    assert "# not a heading" in parsed.sections[0].text


def test_docx_headings_and_tables():
    d = docx.Document()
    d.add_heading("Quarterly Report", level=1)
    d.add_paragraph("Revenue was 12 million.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Region", "Sales"
    table.cell(1, 0).text, table.cell(1, 1).text = "EMEA", "5M"
    buf = io.BytesIO()
    d.save(buf)
    parsed = parse(buf.getvalue(), DOCX)
    assert parsed.sections[0].heading == "Quarterly Report"
    assert "12 million" in parsed.sections[0].text
    assert any("EMEA | 5M" in s.text for s in parsed.sections)


def test_plain_text_utf16():
    parsed = parse("héllo wörld".encode("utf-16"), TEXT)
    assert parsed.sections[0].text == "héllo wörld"


@pytest.mark.parametrize(
    "filename,head",
    [("report.pdf", b"PK\x03\x04"), ("notes.docx", b"%PDF-1.7"), ("a.txt", b"ab\x00cd"), ("x.exe", b"MZ")],
)
def test_detect_media_type_rejects_mismatches(filename, head):
    with pytest.raises(UnsupportedFile):
        detect_media_type(filename, head)


def test_detect_media_type_accepts_valid():
    assert detect_media_type("A.PDF", b"%PDF-1.4") == PDF
    assert detect_media_type("notes.md", b"# Title") == MARKDOWN


def test_corrupt_pdf_is_unsupported():
    with pytest.raises(UnsupportedFile):
        parse(b"%PDF-1.4 garbage", PDF)
